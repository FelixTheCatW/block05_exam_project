import shutil
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference, ScatterChart
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.marker import Marker
from openpyxl.chart.series_factory import SeriesFactory
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.worksheet.worksheet import Worksheet

WS_ANCHOR = "E11"

SCATTER_BG = "0E0E0E"        # фон диаграммы (почти чёрный)
SCATTER_PLOT_BG = "171B1F"   # фон области построения
SCATTER_GRID = "39414A"      # линии сетки
SCATTER_TEXT = "D7DEE5"      # текст осей и заголовка
SCATTER_POINT = "00E5FF"     # цвет точек (неоновый циан)


def _style(title: str, chart) -> None:
    chart.style = 13

    chart.title = title
    chart.title.overlay = False

    chart.legend.position = "b"  # 'b' - снизу, 'r' - справа, 't' - сверху, 'l' - слева
    chart.legend.overlay = False

    if chart.y_axis is not None and chart.y_axis.title is not None:
        chart.y_axis.title.overlay = False

    if chart.x_axis is not None and chart.x_axis.title is not None:
        chart.x_axis.title.overlay = False


def set_axis_title(axis, title: str, font_size: int = 1400) -> None:
    axis.title = title
    axis.title.overlay = False


def chart_by_weekday(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=8)
    cats = Reference(ws, min_col=6, min_row=2, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Ккал"
    chart.x_axis.title = "День недели"

    _style("Средние и медианные калории по дням недели", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_by_category(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    data = Reference(ws, min_col=3, max_col=5, min_row=1, max_row=4)
    cats = Reference(ws, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Граммы"
    chart.x_axis.title = "Категория"
    _style("Средние макронутриенты по категориям", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_category_pie(
    ws: Worksheet,
    n: int,
) -> None:
    pie = PieChart()
    # pie.title = "Доля записей по категориям калорийности"
    pie.style = 13

    pie.title = "Доля записей по категориям калорийности"
    pie.title.overlay = False

    pie.legend.position = "b"  # 'b' - снизу, 'r' - справа, 't' - сверху, 'l' - слева
    pie.legend.overlay = False

    data = Reference(ws, min_col=9, min_row=1, max_row=4)
    cats = Reference(ws, min_col=1, min_row=2, max_row=4)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(cats)
    pie.height = 8
    pie.width = 12    
    ws.add_chart(pie, "N4")


def chart_norm_cat(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    data = Reference(ws, min_col=3, max_col=5, min_row=1, max_row=4)
    cats = Reference(ws, min_col=1, min_row=2, max_row=4)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "% калорий"
    _style("Доля калорий из макронутриентов по категориям", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_norm_weekday(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "percentStacked"
    chart.overlap = 100
    data = Reference(ws, min_col=3, max_col=5, min_row=1, max_row=8)
    cats = Reference(ws, min_col=7, min_row=2, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "% калорий"
    _style("Доля калорий из макронутриентов по дням недели", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_pivot_day_cat(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    data = Reference(ws, min_col=1, max_col=3, min_row=1, max_row=8)
    cats = Reference(ws, min_col=1, min_row=3, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Средние ккал"
    _style("Средние калории: день недели x категория", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_pivot_day_protein(
    ws: Worksheet,
    n: int,
) -> None:
    chart = BarChart()
    chart.type = "col"
    data = Reference(ws, min_col=2, max_col=3, min_row=1, max_row=8)
    cats = Reference(ws, min_col=1, min_row=2, max_row=8)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Число записей"
    _style("Число записей: обычные vs высокий белок", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_calorie_distribution(
    ws: Worksheet,
    n: int,
) -> None:
    header_row = 14
    data_first = header_row + 1
    data_last = header_row + 6
    chart = BarChart()
    chart.type = "col"
    data = Reference(ws, min_col=2, min_row=header_row, max_row=data_last)
    cats = Reference(ws, min_col=1, min_row=data_first, max_row=data_last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "Записей"
    chart.x_axis.title = "Ккал"
    _style("Распределение записей по калорийности", chart)
    ws.add_chart(chart, WS_ANCHOR)


def chart_scatter_calories_protein(ws: Worksheet, n: int) -> None:
    headers = {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value in ("calories", "protein")
    }
    if not {"calories", "protein"} <= set(headers):
        return
    chart = ScatterChart()
    chart.title = "Калории x белок (первые 200 записей)"
    xvalues = Reference(ws, min_col=headers["calories"], min_row=2, max_row=n + 1)
    yvalues = Reference(ws, min_col=headers["protein"], min_row=1, max_row=n + 1)
    series = SeriesFactory(yvalues, xvalues=xvalues, title_from_data=True)
    series.marker = Marker(symbol="circle", size=3)
    series.marker.spPr = GraphicalProperties(
        solidFill=SCATTER_POINT,
        ln=LineProperties(solidFill=SCATTER_POINT, w=22700),
    )
    series.spPr = GraphicalProperties(ln=LineProperties(noFill=True))
    chart.series.append(series)
    chart.x_axis.title = "Калории"
    chart.y_axis.title = "Белок, граммы"
    grid = ChartLines(
        spPr=GraphicalProperties(ln=LineProperties(solidFill=SCATTER_GRID, w=9525))
    )
    chart.x_axis.majorGridlines = grid
    chart.y_axis.majorGridlines = grid
    _style("Калории x белок (первые 200 записей)", chart)
    chart.legend = None
    ws.add_chart(chart, WS_ANCHOR)


def write_calorie_bins(
    writer: pd.ExcelWriter,
    sheet_name: str,
    df: pd.DataFrame,
    start_row: int,
) -> None:
    labels = ["0-500", "500-1000", "1000-1500", "1500-2000", "2000-2500", "2500-3000"]
    cal_range = pd.cut(
        df["calories"],
        bins=[0, 500, 1000, 1500, 2000, 2500, 3000],
        labels=labels,
        right=False,
    )
    dist = cal_range.value_counts().reindex(labels).reset_index()
    dist.columns = ["calories_range", "count"]
    dist.to_excel(writer, sheet_name=sheet_name, startrow=13, index=False)


_C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"

ET.register_namespace("c", _C_NS)
ET.register_namespace("a", _A_NS)


def _a_el(tag: str, **attrs) -> ET.Element:
    el = ET.Element(f"{{{_A_NS}}}{tag}")
    for key, value in attrs.items():
        el.set(key, str(value))
    return el


def _a_solid(color: str) -> ET.Element:
    fill = _a_el("solidFill")
    fill.append(_a_el("srgbClr", val=color))
    return fill


def _a_ln(color: str, w: int | None = None) -> ET.Element:
    ln = _a_el("ln")
    ln.append(_a_solid(color))
    if w is not None:
        ln.set("w", str(w))
    return ln


def _style_scatter_root(root: ET.Element) -> None:
    """Тёмный фон, светлый текст, сетка и 'свечение' точек для scatter-чарта."""
    chart = root.find(f"{{{_C_NS}}}chart")
    if chart is None:
        return

    # 1. Фон диаграммы (chartSpace) — сразу после <c:chart>
    bg = ET.Element(f"{{{_C_NS}}}spPr")
    bg.append(_a_solid(SCATTER_BG))
    root.insert(list(root).index(chart) + 1, bg)

    # 2. Фон области построения (plotArea)
    plot_area = chart.find(f"{{{_C_NS}}}plotArea")
    if plot_area is not None:
        plot_bg = ET.Element(f"{{{_C_NS}}}spPr")
        plot_bg.append(_a_solid(SCATTER_PLOT_BG))
        plot_area.append(plot_bg)

    # 3. Оси: серые линии и светлый текст подписей
    for tag in ("valAx", "catAx"):
        for axis in root.iter(f"{{{_C_NS}}}{tag}"):
            cross = axis.find(f"{{{_C_NS}}}crossAx")
            where = list(axis).index(cross) if cross is not None else len(axis)
            if axis.find(f"{{{_C_NS}}}spPr") is None:
                axis_sp = ET.Element(f"{{{_C_NS}}}spPr")
                axis_sp.append(_a_ln(SCATTER_GRID, w=12700))
                axis.insert(where, axis_sp)
                where += 1
            if axis.find(f"{{{_C_NS}}}txPr") is None:
                tx_pr = ET.Element(f"{{{_C_NS}}}txPr")
                tx_pr.append(_a_el("bodyPr"))
                tx_pr.append(_a_el("lstStyle"))
                paragraph = _a_el("p")
                p_pr = _a_el("pPr")
                def_rpr = _a_el("defRPr")
                def_rpr.append(_a_solid(SCATTER_TEXT))
                p_pr.append(def_rpr)
                paragraph.append(p_pr)
                paragraph.append(_a_el("endParaRPr"))
                tx_pr.append(paragraph)
                axis.insert(where, tx_pr)

    # 4. Свечение точек (effectLst/glow) у всех маркеров
    for marker in root.iter(f"{{{_C_NS}}}marker"):
        sp_pr = marker.find(f"{{{_C_NS}}}spPr")
        if sp_pr is None:
            continue
        effects = _a_el("effectLst")
        glow = _a_el("glow")
        colour = _a_el("srgbClr", val=SCATTER_POINT)
        colour.append(_a_el("alpha", val="90000"))
        glow.append(colour)
        glow.append(_a_el("rad", val="100000"))
        effects.append(glow)
        sp_pr.append(effects)

    # 5. Светлый цвет заголовка
    title = chart.find(f"{{{_C_NS}}}title")
    if title is not None:
        for def_rpr in title.iter(f"{{{_A_NS}}}defRPr"):
            if def_rpr.find(f"{{{_A_NS}}}solidFill") is None:
                def_rpr.append(_a_solid(SCATTER_TEXT))


def style_scatter_chart(path: str | Path) -> None:
    """
    Патчит XML точечных (scatter) диаграмм в сохранённом xlsx:
    тёмный фон, свечение маркеров, светлый текст. Остальные чарты не трогает.
    """
    path = Path(path)
    with zipfile.ZipFile(path) as zin:
        infos = zin.infolist()
        payload = {item.filename: zin.read(item.filename) for item in infos}

    changed = False
    for name, data in payload.items():
        if not name.startswith("xl/charts/") or b"<scatterChart>" not in data:
            continue
        root = ET.fromstring(data)
        _style_scatter_root(root)
        xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            + ET.tostring(root, encoding="unicode")
        )
        payload[name] = xml.encode("utf-8")
        changed = True

    if not changed:
        return

    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in infos:
            zout.writestr(item, payload[item.filename])
    path.unlink()
    shutil.move(str(tmp), str(path))
