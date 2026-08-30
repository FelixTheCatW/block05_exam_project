from pathlib import Path

import pandas as pd

import src.excel_charts as chart


def export_to_excel(
    df: pd.DataFrame,
    stats: pd.DataFrame,
    by_weekday: pd.DataFrame,
    by_category: pd.DataFrame,
    pivot_cat: pd.DataFrame,
    pivot_protein: pd.DataFrame,
    corr: pd.DataFrame,
    norm_cat: pd.DataFrame,
    norm_weekday: pd.DataFrame,
    output_path: Path,
) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        write_and_style_table(writer, stats, "Статистика")
        chart.write_calorie_bins(writer, "Статистика", df, stats.shape[0] + 3)

        write_and_style_table(writer, by_weekday, "По дням недели")
        # write_data_and_chart(writer, "По дням недели", by_weekday)
        write_and_style_table(writer, by_category, "По категориям")
        write_and_style_table(writer, norm_cat, "Нормализация категорий")
        write_and_style_table(writer, norm_weekday, "Распред. по дням")
        write_and_style_table(writer, pivot_cat, "Сводная по категориям")
        write_and_style_table(writer, pivot_protein, "Белок")
        write_and_style_table(writer, corr, "Корреляция")

        chart.chart_by_weekday(writer.sheets["По дням недели"], by_weekday.shape[0])
        chart.chart_by_category(writer.sheets["По категориям"], by_category.shape[0])
        chart.chart_category_pie(writer.sheets["По категориям"], by_category.shape[0])
        chart.chart_norm_cat(writer.sheets["Нормализация категорий"], norm_cat.shape[0])
        chart.chart_norm_weekday(
            writer.sheets["Распред. по дням"], norm_weekday.shape[0]
        )
        chart.chart_pivot_day_cat(
            writer.sheets["Сводная по категориям"], pivot_cat.shape[0]
        )
        chart.chart_pivot_day_protein(writer.sheets["Белок"], pivot_protein.shape[0])
        chart.chart_calorie_distribution(writer.sheets["Статистика"], 6)
        
        exclude_cols = ['fiber', 'potass.', 'iron', 'calcium', 'sat fat', 'chol', 'vit a', 'vit c', 'trn fat', 'mon fat', 'ply fat']
        sample_200 = df.head(200).drop(columns=[col for col in exclude_cols if col in df.columns], errors='ignore')
        sample_200.to_excel(writer, sheet_name="Первые 200")
        chart.chart_scatter_calories_protein(
            writer.sheets["Первые 200"], sample_200.shape[0]
        )

    chart.style_scatter_chart(output_path)

    print(f"Отчёт с диаграммами сохранён:\n{output_path}")


def write_data_and_chart(
    writer: pd.ExcelWriter,
    df: pd.DataFrame,
    ws_title: str,
):
    ws = write_and_style_table(writer, df, ws_title)
    chart.chart_by_weekday(ws, df.shape[0])


import uuid

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def write_and_style_table(
    writer, df, sheet_name, index=False, style_name="TableStyleMedium4"
):
    df.to_excel(writer, sheet_name=sheet_name, index=index)
    ws = writer.sheets[sheet_name]

    if df.empty:
        return ws

    rows = df.shape[0]
    cols = df.shape[1]
    total_cols = cols + (1 if index else 0)
    max_row = rows + 1
    max_col = total_cols
    start_cell = "A1"
    end_cell = f"{get_column_letter(max_col)}{max_row}"
    table_range = f"{start_cell}:{end_cell}"

    for tbl in list(ws.tables.values()):
        ws._tables.remove(tbl)

    safe_sheet_name = sheet_name.replace(" ", "_").replace("(", "").replace(")", "")
    table_name = f"Table_{safe_sheet_name}_{uuid.uuid4().hex[:6]}"

    try:
        table = Table(displayName=table_name, ref=table_range)
        style = TableStyleInfo(
            name=style_name,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)
    except Exception as e:
        print(f"Не удалось применить стиль к листу '{sheet_name}': {e}")

    return ws
