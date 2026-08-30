import uuid
from functools import partial
from pathlib import Path

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

import src.excel_charts as chart
from src.utils.pipe import Pipe


class ReportWriter:
    def __init__(self, output_path: Path):
        self.output_path = output_path
        self.writer = None
        self.pipe = None

    def __enter__(self):
        self.writer = pd.ExcelWriter(self.output_path, engine="openpyxl")
        self.pipe = Pipe(self.writer)
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):

        if self.writer:
            self.writer.close()

        # chart.style_scatter_chart(self.output_path)
        print(f"Отчёт с диаграммами сохранён:\n{self.output_path}")

        return False


def write_table(df, sheet_name, writer):
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]

    if df.empty:
        return ws

    max_row = df.shape[0] + 1
    max_col = df.shape[1]
    start_cell = "A1"
    end_cell = f"{get_column_letter(max_col)}{max_row}"
    table_range = f"{start_cell}:{end_cell}"

    safe_sheet_name = sheet_name.replace(" ", "_").replace("(", "").replace(")", "")
    table_name = f"Table_{safe_sheet_name}_{uuid.uuid4().hex[:6]}"

    try:
        table = Table(displayName=table_name, ref=table_range)
        style = TableStyleInfo(
            name="TableStyleMedium4",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        table.tableStyleInfo = style
        ws.add_table(table)
    except Exception as e:
        print(f"Не удалось применить стиль к листу '{sheet_name}': {e}")

    return ws
